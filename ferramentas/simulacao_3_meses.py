"""
Simulação de TRÊS MESES de uso do backup, em minutos, sem tocar em nada real.

Roda o programa de verdade (backup.py, backup_comprovantes.py e a rodada do notebook,
ferramentas/backup_no_servidor.py) contra uma planilha, um Drive, um servidor e um relógio
DE MENTIRA, todos dentro de uma pasta temporária. Nenhuma chamada sai para a internet.

O que acontece nos três meses (01/10/2026 a 05/01/2027):
  - uso normal: cobranças novas, pagamentos (com e sem comprovante), observações,
    renegociação de parcela, cobrança excluída e restaurada;
  - o notebook roda 10h30 e 16h30 nos dias úteis; a nuvem roda 6x ao dia com atraso e,
    às vezes, pula a rodada (como o GitHub faz de verdade);
  - imprevistos: duas mudanças no mesmo minuto, rodada interrompida no meio, Drive fora do
    ar, notebook desligado nas férias, GitHub fora do ar, servidor inalcançável, alguém
    apagando cópia do servidor e comprovante do Drive, planilha que ENCOLHE, planilha que
    não abre, arquivo do servidor estragado, virada de mês e virada de ano.

No fim confere as regras (nenhuma cópia reescrita, nenhuma duplicada, tudo na pasta do mês
certo, servidor igual ao notebook, cada comprovante baixado uma única vez, alarmes só quando
devem tocar) e mostra os números.

Rodar:  py ferramentas\\simulacao_3_meses.py          (--manter guarda a pasta para olhar)
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import os
import random
import re
import shutil
import sys
import tempfile
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
TMP = Path(tempfile.mkdtemp(prefix="sim3m_"))
NOTEBOOK = TMP / "notebook"
REDE, REDE_FORA = TMP / "rede", TMP / "rede_fora_do_ar"
SERVIDOR = REDE / "Pedro" / "BACKUP" / "BACKUP COBRANCAS"
COFRE = TMP / "cofre_github"
SERVIDOR.parent.mkdir(parents=True)

# Tudo que a rodada do notebook lê na partida aponta para a pasta temporária
(TMP / "config.json").write_text(json.dumps({
    "pasta_notebook": str(NOTEBOOK), "pasta_servidor": str(SERVIDOR), "comprovantes_drive": "RAIZ_COMPROVANTES"}),
    encoding="utf-8")
(TMP / "secrets.toml").write_text(
    'spreadsheet_id = "simulada"\ndrive_oauth_client_id = "x"\ndrive_oauth_client_secret = "x"\n'
    'drive_oauth_refresh_token = "x"\n[gcp_service_account]\ntype = "simulada"\n', encoding="utf-8")
os.environ["COBRANCA_BACKUP_CONFIG"] = str(TMP / "config.json")
os.environ["COBRANCA_BACKUP_SECRETS"] = str(TMP / "secrets.toml")
os.environ["DRIVE_OAUTH"] = '{"client_id": "x", "client_secret": "x", "refresh_token": "x"}'

sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "ferramentas"))
import backup  # noqa: E402
import backup_comprovantes as bc  # noqa: E402
import backup_no_servidor as notebook  # noqa: E402
import dados  # noqa: E402
import nucleo  # noqa: E402

TZ = backup.TIMEZONE
INICIO, FIM = date(2026, 10, 1), date(2027, 1, 5)
rng = random.Random(2026)


class Relogio:
    agora = datetime(2026, 10, 1, 8, 0, tzinfo=TZ)


def dt(dia: date, h: int, m: int, s: int = 0) -> datetime:
    return datetime(dia.year, dia.month, dia.day, h, m, s, tzinfo=TZ)


# ── Drive de mentira ──────────────────────────────────────────────────────────

class _Chamada:
    def __init__(self, f):
        self._f = f

    def execute(self):
        return self._f()


class DriveFalso:
    """O pedaço da API do Drive que o backup usa, guardado na memória."""

    def __init__(self):
        self.arqs, self._n, self.fora = {}, 0, False
        self.baixados, self.reescritos = [], []
        self.arqs["RAIZ_COMPROVANTES"] = self._novo({"name": "Cobranças Vithall - Comprovantes",
                                                     "mimeType": backup._MIME_PASTA}, b"", "RAIZ_COMPROVANTES")

    def files(self):
        return self

    def _novo(self, body, conteudo, fid=None):
        self._n += 1
        agora = Relogio.agora.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
        return {"id": fid or f"arq{self._n:05d}", "name": body["name"],
                "mimeType": body.get("mimeType") or "application/octet-stream",
                "parents": list(body.get("parents") or ["root"]),
                "appProperties": dict(body.get("appProperties") or {}),
                "createdTime": agora, "modifiedTime": agora, "conteudo": conteudo}

    def _vista(self, a):
        v = {k: a[k] for k in ("id", "name", "mimeType", "parents", "appProperties", "createdTime", "modifiedTime")}
        if a["mimeType"] != backup._MIME_PASTA:
            v.update(size=str(len(a["conteudo"])), md5Checksum=hashlib.md5(a["conteudo"]).hexdigest(),
                     webViewLink=f"https://drive.google.com/file/d/{a['id']}/view?usp=drivesdk")
        return v

    def _vivo(self):
        if self.fora:
            raise OSError("Drive fora do ar (simulado)")

    @staticmethod
    def _casa(a, q):
        solto = lambda s: s.replace("\\'", "'").replace("\\\\", "\\")
        m = re.search(r"appProperties has \{ key='([^']*)' and value='([^']*)' \}", q)
        if m:
            if a["appProperties"].get(m[1]) != m[2]:
                return False
            q = q.replace(m[0], "")
        m = re.search(r"\(name = '(.*?)' or name = '(.*?)'\)", q)
        if m:
            if a["name"] not in (solto(m[1]), solto(m[2])):
                return False
            q = q.replace(m[0], "")
        for c in (p.strip() for p in q.split(" and ")):
            if c in ("", "trashed = false"):
                continue
            if m := re.fullmatch(r"name = '(.*)'", c):
                ok = a["name"] == solto(m[1])
            elif m := re.fullmatch(r"'(.*)' in parents", c):
                ok = m[1] in a["parents"]
            elif m := re.fullmatch(r"mimeType = '(.*)'", c):
                ok = a["mimeType"] == m[1]
            elif m := re.fullmatch(r"mimeType != '(.*)'", c):
                ok = a["mimeType"] != m[1]
            else:
                raise ValueError("consulta que o Drive de mentira não entende: " + c)
            if not ok:
                return False
        return True

    def list(self, q="", orderBy=None, pageSize=100, pageToken=None, fields=None):
        def f():
            self._vivo()
            achados = [a for a in self.arqs.values() if self._casa(a, q)]
            if orderBy:
                campo, _, sentido = orderBy.partition(" ")
                achados.sort(key=lambda a: a[campo], reverse=(sentido == "desc"))
            ini = int(pageToken or 0)
            r = {"files": [self._vista(a) for a in achados[ini:ini + pageSize]]}
            if ini + pageSize < len(achados):
                r["nextPageToken"] = str(ini + pageSize)
            return r
        return _Chamada(f)

    def create(self, body, media_body=None, fields=None):
        def f():
            self._vivo()
            conteudo = media_body.getbytes(0, media_body.size()) if media_body is not None else b""
            a = self._novo(body, conteudo)
            self.arqs[a["id"]] = a
            return self._vista(a)
        return _Chamada(f)

    def update(self, fileId, body=None, media_body=None, addParents=None, removeParents=None, fields=None):
        def f():
            self._vivo()
            a = self.arqs[fileId]
            if media_body is not None:
                if backup.SUFIXO in a["name"]:
                    self.reescritos.append(a["name"])
                a["conteudo"] = media_body.getbytes(0, media_body.size())
                a["modifiedTime"] = Relogio.agora.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
            if body and "name" in body:
                a["name"] = body["name"]
            if removeParents:
                a["parents"] = [p for p in a["parents"] if p not in removeParents.split(",")]
            if addParents:
                a["parents"] += addParents.split(",")
            return self._vista(a)
        return _Chamada(f)

    def get(self, fileId, fields=None):
        return _Chamada(lambda: (self._vivo(), self._vista(self.arqs[fileId]))[1])

    def get_media(self, fileId):
        def f():
            self._vivo()
            self.baixados.append(fileId)
            return self.arqs[fileId]["conteudo"]
        return _Chamada(f)

    def delete(self, fileId):
        return _Chamada(lambda: self.arqs.pop(fileId) and None)


DRIVE = DriveFalso()

# ── Planilha de mentira (o mesmo "banco local" que o app usa para testar telas) ───────────

BANCO = dados.Local(str(TMP / "planilha.local.json"))
ESTADO_SIM = {"leitura_fora": False}
SEC_APP = {"drive_folder_id": "RAIZ_COMPROVANTES", "drive_oauth_client_id": "x",
           "drive_oauth_client_secret": "x", "drive_oauth_refresh_token": "x"}


def abas_da_planilha() -> dict:
    d = BANCO._ler()
    return {
        "Página1": [],
        "_Cobrancas": [nucleo.COBRANCAS_HEADERS] + [[c.get(h, "") for h in nucleo.COBRANCAS_HEADERS] for c in d["cobrancas"]],
        "_Parcelas": [nucleo.PARCELAS_HEADERS] + [[p.get(h, "") for h in nucleo.PARCELAS_HEADERS] for p in d["parcelas"]],
        "_Config": [dados.CONFIG_HEADERS] + [[k, v] for k, v in d.get("config", {}).items()],
    }


def ler_planilha_falsa(_id=None):
    if ESTADO_SIM["leitura_fora"]:
        raise RuntimeError("o Google não respondeu (simulado)")
    return backup.montar("Cobranças Vithall (simulada)", abas_da_planilha())


AVISOS = []   # (quando, texto) das janelas de aviso que apareceriam na tela
backup.agora_sp = lambda: Relogio.agora
backup.ler_planilha = ler_planilha_falsa
backup.servico_drive = lambda _cred: DRIVE
dados._drive = lambda _sec: DRIVE
notebook._avisar_na_tela = lambda texto: AVISOS.append((Relogio.agora, texto))

# ── O uso do app ──────────────────────────────────────────────────────────────

NOMES = ["Ana", "Bruno", "Carla", "Diego", "Elisa", "Fábio", "Gisele", "Hugo", "Iara", "Jonas", "Karen", "Lucas",
         "Marta", "Nilo", "Olga", "Paulo", "Rita", "Saulo", "Tânia", "Vitor"]
SOBRENOMES = ["Souza", "Lima", "Becker", "Schmidt", "D'Ávila", "Rocha", "Müller", "Kunz", "Alves", "Finger"]
PESSOAS = ["Gabi", "Ana", "Pedro"]
CONTAGEM = {"cobrancas": 0, "pagamentos": 0, "comprovantes": 0, "observacoes": 0, "renegociacoes": 0,
            "exclusoes": 0, "restauracoes": 0}
COMPROVANTES_ENVIADOS = {}   # id no Drive -> md5 do conteúdo original


def nova_cobranca(hoje: date):
    tipo = rng.choice([nucleo.TIPO_PIX, nucleo.TIPO_PIX, nucleo.TIPO_CARTAO])
    turma = rng.choice("LLLVIP") + str(rng.randint(320, 360))
    valor, n = float(rng.choice([250, 500, 750, 1000, 1900])), rng.randint(2, 10)
    pessoa = rng.choice(PESSOAS)
    parcelas = []
    if rng.random() < 0.6:
        parcelas.append({"Nº": 0, "Vencimento": hoje, "Valor": valor, "Status": nucleo.STATUS_PAGA,
                         "Pago em": hoje, "Marcado por": pessoa})
    primeira = hoje + timedelta(days=rng.randint(2, 25))
    for i in range(1, n + 1):
        p = {"Nº": i, "Vencimento": nucleo.mes_mais(primeira, i - 1), "Valor": valor}
        if tipo == nucleo.TIPO_CARTAO:
            p["Vezes no cartão"] = rng.randint(1, 10)
        parcelas.append(p)
    BANCO.criar_cobranca({
        "Tipo": tipo, "Cliente": f"{rng.choice(NOMES)} {rng.choice(SOBRENOMES)}", "Turma": turma,
        "Treinamento": nucleo.treinamento_da_turma(turma), "Valor Total": valor * len(parcelas),
        "Criada em": hoje, "Criada por": pessoa, "Cidade": rng.choice(["", "", "Lajeado", "SCS", "Caxias"]),
        "Observações": ""}, parcelas)
    CONTAGEM["cobrancas"] += 1


def comprovante_falso() -> bytes:
    return b"%PDF-1.4 comprovante " + rng.randbytes(rng.randint(60_000, 180_000))


def usar_o_app(agora: datetime, intensidade: float = 1.0):
    hoje = agora.date()
    if rng.random() < 0.30 * intensidade:
        nova_cobranca(hoje)
    d = BANCO._ler()
    cobs = {c["ID"]: c for c in d["cobrancas"]}
    ativas = [c for c in d["cobrancas"] if c.get("Situação") != nucleo.SITUACAO_EXCLUIDA]
    for p in d["parcelas"]:
        c = cobs.get(p["ID Cobrança"])
        venc = nucleo.parse_data(p.get("Vencimento"))
        if (not c or c.get("Situação") == nucleo.SITUACAO_EXCLUIDA or p.get("Status") == nucleo.STATUS_PAGA
                or venc is None or venc > hoje + timedelta(days=2) or rng.random() > 0.22 * intensidade):
            continue
        campos = {"Status": nucleo.STATUS_PAGA, "Pago em": hoje, "Marcado por": rng.choice(PESSOAS)}
        if rng.random() < 0.65:
            conteudo = comprovante_falso()
            # Às vezes a pessoa lança no dia seguinte o pagamento de ontem (data do nome != dia do envio)
            quando = hoje - timedelta(days=1) if rng.random() < 0.15 else hoje
            link = dados.enviar_comprovante(SEC_APP, conteudo, "comprovante.pdf", c["Treinamento"],
                                            f"{c['Cliente']} - {c['Turma']} - Parcela {p['Nº']}", quando)
            campos["Comprovante"] = link
            COMPROVANTES_ENVIADOS[re.search(r"/d/([^/]+)", link).group(1)] = hashlib.md5(conteudo).hexdigest()
            CONTAGEM["comprovantes"] += 1
        BANCO.atualizar_parcela(p["ID Cobrança"], int(p["Nº"]), campos)
        CONTAGEM["pagamentos"] += 1
    if ativas and rng.random() < 0.06 * intensidade:
        BANCO.atualizar_cobranca(rng.choice(ativas)["ID"], {"Observações": rng.choice(
            ["pago em permuta", "=combinado por telefone", "vai quitar em dezembro", "trocou o cartão"])})
        CONTAGEM["observacoes"] += 1
    if ativas and rng.random() < 0.04 * intensidade:   # renegociação: sai a última parcela aberta, entra outra
        c = rng.choice(ativas)
        abertas = [p for p in BANCO._ler()["parcelas"] if p["ID Cobrança"] == c["ID"] and p.get("Status") != nucleo.STATUS_PAGA]
        if abertas:
            ultima = max(abertas, key=lambda p: int(p["Nº"]))
            BANCO.remover_parcela(c["ID"], int(ultima["Nº"]))
            BANCO.adicionar_parcela(c["ID"], {"Nº": int(ultima["Nº"]) + 1, "Valor": nucleo.valor_para_float(ultima["Valor"]),
                                              "Vencimento": hoje + timedelta(days=40)})
            CONTAGEM["renegociacoes"] += 1
    if ativas and rng.random() < 0.015 * intensidade:
        dados.excluir_cobranca(BANCO, rng.choice(ativas)["ID"], hoje)
        CONTAGEM["exclusoes"] += 1
    excluidas = [c for c in BANCO._ler()["cobrancas"] if c.get("Situação") == nucleo.SITUACAO_EXCLUIDA]
    if excluidas and rng.random() < 0.01 * intensidade:
        dados.restaurar_cobranca(BANCO, rng.choice(excluidas)["ID"])
        CONTAGEM["restauracoes"] += 1


# ── Acompanhamento ────────────────────────────────────────────────────────────

DESTINOS = ("notebook", "servidor", "drive", "cofre")
VISTOS = {d: set() for d in DESTINOS}          # impressões que cada destino tem que ter guardado
PENDENTE = {d: None for d in DESTINOS + ("qualquer",)}   # hora da mudança mais antiga ainda sem cópia
ESPERAS = {d: [] for d in DESTINOS + ("qualquer",)}
PRIMEIRA_VEZ = {}                               # arquivo datado -> sha256 quando apareceu
RODADAS = {"notebook": 0, "nuvem": 0, "nuvem_com_erro": [], "notebook_tempo": []}
ULTIMA_IMPRESSAO = [None]
ESTRAGADOS = set()                              # arquivos que a simulação estraga de propósito


def houve_mudanca():
    atual = backup.impressao_digital(abas_da_planilha())
    if atual != ULTIMA_IMPRESSAO[0]:
        ULTIMA_IMPRESSAO[0] = atual
        for d in PENDENTE:
            PENDENTE[d] = PENDENTE[d] or Relogio.agora


def guardou(destino: str):
    VISTOS[destino].add(backup.impressao_digital(abas_da_planilha()))
    for d in (destino, "qualquer"):
        if PENDENTE[d]:
            ESPERAS[d].append((Relogio.agora - PENDENTE[d]).total_seconds() / 3600)
            PENDENTE[d] = None


def anotar_arquivos_novos():
    for raiz in (NOTEBOOK / "copias", SERVIDOR, COFRE):
        if raiz.exists():
            for arq in raiz.rglob("*"):
                if arq.is_file() and arq.name not in (backup.MARCA, "LEIA-ME.txt", bc.INDICE, bc.LISTA) \
                        and not arq.name.endswith(".parcial") and str(arq) not in PRIMEIRA_VEZ:
                    PRIMEIRA_VEZ[str(arq)] = hashlib.sha256(arq.read_bytes()).hexdigest()


def rodar_notebook():
    t0 = time.perf_counter()
    with contextlib.redirect_stdout(io.StringIO()):
        notebook.main()
    RODADAS["notebook_tempo"].append(time.perf_counter() - t0)
    RODADAS["notebook"] += 1
    estado = json.loads((NOTEBOOK / "estado.json").read_text(encoding="utf-8"))
    agora = Relogio.agora.isoformat(timespec="seconds")
    for chave, destino in (("leitura_ok", "notebook"), ("servidor_ok", "servidor"), ("drive_ok", "drive")):
        # o servidor só "tem" o estado de agora se o notebook também o leu nesta rodada
        if estado.get(chave) == agora and estado.get("leitura_ok") == agora:
            guardou(destino)
    anotar_arquivos_novos()


def rodar_nuvem():
    RODADAS["nuvem"] += 1
    try:
        copia = backup.ler_planilha("simulada")
    except Exception as e:
        RODADAS["nuvem_com_erro"].append((Relogio.agora, f"não leu a planilha: {e}"))
        return
    erros = backup.executar(copia, pasta=COFRE, drive=True, saida=lambda *_: None)
    if not any(e.startswith("pasta") for e in erros):
        guardou("cofre")
    if not any(e.startswith("Drive") for e in erros):
        guardou("drive")
    if erros:
        RODADAS["nuvem_com_erro"].append((Relogio.agora, "; ".join(erros)))
    anotar_arquivos_novos()


def servidor_no_ar(sim: bool):
    if sim and REDE_FORA.exists():
        REDE_FORA.rename(REDE)
    elif not sim and REDE.exists():
        REDE.rename(REDE_FORA)


# ── Imprevistos ───────────────────────────────────────────────────────────────

FERIADOS = {date(2026, 10, 12), date(2026, 11, 2), date(2026, 12, 25), date(2027, 1, 1)}
entre = lambda dia, a, b: a <= dia <= b
DRIVE_FORA = (date(2026, 11, 3), date(2026, 11, 6))
FERIAS = (date(2026, 11, 10), date(2026, 11, 19))           # notebook desligado
GITHUB_FORA = (date(2026, 11, 25), date(2026, 11, 29))
SERVIDOR_FORA = (date(2026, 12, 1), date(2026, 12, 6))
LEITURA_FORA = (date(2026, 12, 15), date(2026, 12, 18))
RECESSO = (date(2026, 12, 22), date(2027, 1, 2))            # quase ninguém usa o app
NOTAS = []                                                   # o que cada imprevisto mostrou


def mesmo_minuto():
    """Duas rodadas e uma mudança dentro do mesmo minuto: a segunda cópia tem que ganhar o minuto seguinte."""
    base = Relogio.agora
    rodar_notebook()
    Relogio.agora = base + timedelta(seconds=20)
    nova_cobranca(Relogio.agora.date())
    houve_mudanca()
    Relogio.agora = base + timedelta(seconds=40)
    rodar_notebook()
    pasta = Path(NOTEBOOK / "copias", *backup.pastas_da_copia(base.strftime("%Y-%m-%d")))
    nomes = sorted(p.name for p in pasta.glob(base.strftime("%Y-%m-%d") + "*.json"))
    esperado = [(base + timedelta(minutes=k)).strftime("%Y-%m-%d %Hh%M") + backup.SUFIXO + ".json" for k in (0, 1)]
    NOTAS.append(("Duas mudanças no mesmo minuto", nomes[-2:] == esperado,
                  "a segunda cópia ficou com o minuto seguinte" if nomes[-2:] == esperado else f"nomes: {nomes}"))


def rodada_interrompida():
    """Como se o notebook tivesse desligado no meio de uma gravação: sobra de arquivo pela metade."""
    pasta = Path(NOTEBOOK / "copias", *backup.pastas_da_copia("2026-10-27"))
    pasta.mkdir(parents=True, exist_ok=True)
    base = "2026-10-27 16h29" + backup.SUFIXO
    (pasta / (base + ".xlsx.parcial")).write_bytes(b"pela metade")
    (pasta / (base + ".xlsx")).write_bytes(b"xlsx que ficou sem o json")
    cortado = pasta / ("2026-10-27 16h28" + backup.SUFIXO + ".json")
    cortado.write_text('{"abas": {"_Cobr', encoding="utf-8")
    ESTRAGADOS.update({base + ".xlsx", cortado.name})
    antes = RODADAS["notebook"]
    try:
        rodar_notebook()
        ok = True
    except Exception as e:
        ok = False
        NOTAS.append(("Rodada interrompida no meio", False, f"a rodada seguinte quebrou: {e}"))
    if ok:
        NOTAS.append(("Rodada interrompida no meio", RODADAS["notebook"] == antes + 1
                      and not list(SERVIDOR.rglob("*.parcial")),
                      "a rodada seguinte funcionou e o arquivo pela metade não foi para o servidor"))


def alguem_apaga():
    """Somem do servidor uma cópia da planilha e um comprovante; somem do Drive dois comprovantes."""
    copias = sorted(SERVIDOR.rglob(f"*{backup.SUFIXO}.json"))
    alvo = copias[len(copias) // 2]
    alvo.unlink()
    alvo.with_suffix(".xlsx").unlink()
    fotos = sorted((SERVIDOR / bc.PASTA).rglob("*.pdf"))
    fotos[0].unlink()
    ids = [i for i, a in DRIVE.arqs.items() if i in COMPROVANTES_ENVIADOS][:2]
    for i in ids:
        DRIVE.arqs.pop(i)
    rodar_notebook()
    voltou = alvo.exists() and alvo.with_suffix(".xlsx").exists() and fotos[0].exists()
    indice = bc.ler_indice(NOTEBOOK / "copias")["arquivos"]
    ficou = all(i in indice and (NOTEBOOK / "copias" / indice[i]["arquivo"]).exists() for i in ids)
    NOTAS.append(("Alguém apaga cópia e comprovante do servidor", voltou, "a rodada seguinte devolveu os arquivos"))
    NOTAS.append(("Comprovante apagado do Drive", ficou, "continuou guardado no backup"))


ANTES_DE_ENCOLHER = {}


def planilha_encolhe():
    """Três cobranças somem direto da planilha (alguém apagou linhas no Google)."""
    d = BANCO._ler()
    ANTES_DE_ENCOLHER.update(impressao=backup.impressao_digital(abas_da_planilha()), avisos=len(AVISOS),
                             nuvem=len(RODADAS["nuvem_com_erro"]))
    somem = {c["ID"] for c in d["cobrancas"][3:6]}
    d["cobrancas"] = [c for c in d["cobrancas"] if c["ID"] not in somem]
    d["parcelas"] = [p for p in d["parcelas"] if p["ID Cobrança"] not in somem]
    BANCO._gravar(d)
    houve_mudanca()


def planilha_volta():
    """Conserto: reconstrói a planilha a partir da última cópia boa guardada no cofre."""
    alarme_tela = any("encolheu" in t for _, t in AVISOS[ANTES_DE_ENCOLHER["avisos"]:])
    alarme_nuvem = any("encolheu" in t for _, t in RODADAS["nuvem_com_erro"][ANTES_DE_ENCOLHER["nuvem"]:])
    boa = next(c for c in sorted(COFRE.rglob(f"*{backup.SUFIXO}.json"), key=lambda p: p.name, reverse=True)
               if json.loads(c.read_text(encoding="utf-8"))["impressao"] == ANTES_DE_ENCOLHER["impressao"])
    abas = json.loads(boa.read_text(encoding="utf-8"))["abas"]
    linhas = lambda nome: [dict(zip(abas[nome][0], l + [""] * (len(abas[nome][0]) - len(l)))) for l in abas[nome][1:]]
    BANCO._gravar({"cobrancas": linhas("_Cobrancas"), "parcelas": linhas("_Parcelas"),
                   "config": {l[0]: l[1] for l in abas["_Config"][1:]}})
    igual = backup.impressao_digital(abas_da_planilha()) == ANTES_DE_ENCOLHER["impressao"]
    houve_mudanca()
    NOTAS.append(("Planilha encolheu: alarme na tela do notebook", alarme_tela, "a janela de aviso apareceu"))
    NOTAS.append(("Planilha encolheu: alarme no robô da nuvem", alarme_nuvem, "a rodada da nuvem saiu com erro"))
    NOTAS.append(("Planilha reconstruída da última cópia boa", igual, "ficou idêntica ao que era antes de encolher"))


def estraga_no_servidor():
    """Um arquivo de cópia no servidor é alterado (disco ruim, alguém abriu e salvou por cima)."""
    alvo = sorted(SERVIDOR.rglob(f"*{backup.SUFIXO}.xlsx"))[-3]
    alvo.write_bytes(alvo.read_bytes() + b"estragado")
    ESTRAGADOS.add("SERVIDOR:" + alvo.name)


# ── A linha do tempo ──────────────────────────────────────────────────────────

def montar_eventos():
    ev = []
    dia = INICIO
    while dia <= FIM:
        util = dia.weekday() < 5 and dia not in FERIADOS
        if util or (dia.weekday() == 5 and rng.random() < 0.2):
            for h, m in ((9, 5), (11, 20), (14, 10), (15, 45), (17, 40)):
                ev.append((dt(dia, h, m), "uso"))
        ligado = (util or (dia.weekday() == 5 and rng.random() < 0.3)) and not entre(dia, *FERIAS)
        if ligado:
            ev += [(dt(dia, 10, 30), "notebook"), (dt(dia, 16, 30), "notebook")]
        if not entre(dia, *GITHUB_FORA):
            for h, m in ((7, 40), (10, 10), (13, 10), (16, 10), (19, 10), (22, 10)):
                if rng.random() < 0.10:
                    continue                                    # o GitHub pulou esta rodada
                atraso = min(int(rng.expovariate(1 / 50)), 240)  # e atrasa, em média, 50 min
                ev.append((dt(dia, h, m) + timedelta(minutes=atraso), "nuvem"))
        dia += timedelta(days=1)
    ev += [(dt(date(2026, 10, 20), 14, 30), mesmo_minuto), (dt(date(2026, 10, 27), 16, 29, 30), rodada_interrompida),
           (dt(date(2026, 12, 9), 9, 0), alguem_apaga), (dt(date(2026, 12, 11), 11, 0), planilha_encolhe),
           (dt(date(2026, 12, 14), 9, 0), planilha_volta), (dt(date(2026, 12, 21), 12, 0), estraga_no_servidor),
           (dt(FIM, 18, 0), "notebook"), (dt(FIM, 18, 5), "nuvem")]
    return sorted(ev, key=lambda e: e[0])


def simular():
    for chave, valor in (("pessoas", "Pedro, Gabi, Ana"), ("admin", "Pedro"), ("tamanho_texto", "Médio")):
        BANCO.save_config(chave, valor)
    for t, que in montar_eventos():
        Relogio.agora = t
        dia = t.date()
        DRIVE.fora = entre(dia, *DRIVE_FORA)
        ESTADO_SIM["leitura_fora"] = entre(dia, *LEITURA_FORA)
        servidor_no_ar(not entre(dia, *SERVIDOR_FORA))
        if que == "uso":
            # Com o Drive fora o app não sobe comprovante; com a planilha fora ele nem grava
            if not DRIVE.fora and not ESTADO_SIM["leitura_fora"]:
                usar_o_app(t, 0.15 if entre(dia, *RECESSO) else 1.0)
            houve_mudanca()
        elif que == "notebook":
            rodar_notebook()
        elif que == "nuvem":
            rodar_nuvem()
        else:
            que()


# ── Conferência ───────────────────────────────────────────────────────────────

def copias_em(raiz: Path) -> list[tuple[Path, dict | None]]:
    saida = []
    for arq in sorted(raiz.rglob(f"*{backup.SUFIXO}.json"), key=lambda p: p.name):
        try:
            saida.append((arq, json.loads(arq.read_text(encoding="utf-8"))))
        except ValueError:
            saida.append((arq, None))
    return saida


def conferir() -> list[tuple[str, bool, str]]:
    r = list(NOTAS)
    locais = {"notebook": NOTEBOOK / "copias", "servidor": SERVIDOR, "cofre": COFRE}
    impressoes = {}
    for nome, raiz in locais.items():
        copias = [(a, c) for a, c in copias_em(raiz) if a.name not in ESTRAGADOS]
        ruins, fora_do_lugar, repetidas, anterior = [], [], [], None
        for arq, c in copias:
            estragado = "SERVIDOR:" + arq.with_suffix(".xlsx").name in ESTRAGADOS and nome == "servidor"
            if c is None or backup.impressao_digital(c["abas"]) != c["impressao"]:
                ruins.append(arq.name)
            elif not estragado and backup.ler_xlsx(arq.with_suffix(".xlsx").read_bytes()) != c["abas"]:
                ruins.append(arq.with_suffix(".xlsx").name)
            if arq.relative_to(raiz).parts[:-1] != backup.pastas_da_copia(arq.name):
                fora_do_lugar.append(str(arq.relative_to(raiz)))
            if c and anterior and c["impressao"] == anterior:
                repetidas.append(arq.name)
            anterior = c["impressao"] if c else anterior
        impressoes[nome] = {c["impressao"] for _, c in copias if c}
        r.append((f"{nome}: toda cópia abre e o Excel bate com a cópia exata", not ruins, f"{len(copias)} cópias" if not ruins else f"ruins: {ruins[:3]}"))
        r.append((f"{nome}: cada cópia na pasta do seu ano e mês", not fora_do_lugar, "ok" if not fora_do_lugar else str(fora_do_lugar[:3])))
        r.append((f"{nome}: nenhuma cópia repetida à toa", not repetidas, "ok" if not repetidas else str(repetidas[:3])))
    no_drive = sorted((a for a in DRIVE.arqs.values() if a["appProperties"].get("tipo") == "backup-cobrancas"),
                      key=lambda a: a["name"])
    impressoes["drive"] = {a["appProperties"]["impressao"] for a in no_drive}
    nomes = [a["name"] for a in DRIVE.arqs.values() if backup.SUFIXO in a["name"]]
    r.append(("drive: nenhum nome de arquivo repetido", len(nomes) == len(set(nomes)), f"{len(no_drive)} cópias"))
    seguidas = [a["name"] for a, b in zip(no_drive[1:], no_drive) if a["appProperties"]["impressao"] == b["appProperties"]["impressao"]]
    r.append(("drive: nenhuma cópia repetida à toa", not seguidas, "ok" if not seguidas else str(seguidas[:3])))
    r.append(("drive: nenhuma cópia antiga reescrita", not DRIVE.reescritos, "ok" if not DRIVE.reescritos else str(DRIVE.reescritos[:3])))
    for nome in DESTINOS:
        falta = VISTOS[nome] - impressoes[nome]
        r.append((f"{nome}: guardou tudo o que disse que guardou", not falta, f"{len(VISTOS[nome])} estados" if not falta else f"faltam {len(falta)}"))

    mudaram = [Path(p).name for p, sha in PRIMEIRA_VEZ.items()
               if Path(p).exists() and hashlib.sha256(Path(p).read_bytes()).hexdigest() != sha]
    esperados = {n.split(":", 1)[1] for n in ESTRAGADOS if n.startswith("SERVIDOR:")}
    r.append(("Nenhum arquivo guardado foi reescrito pelo backup", set(mudaram) <= esperados,
              "só mudou o que a simulação estragou de propósito" if set(mudaram) <= esperados else str(mudaram[:3])))

    def arvore(raiz):
        return {a.relative_to(raiz).as_posix(): hashlib.sha256(a.read_bytes()).hexdigest()
                for a in raiz.rglob("*") if a.is_file() and not a.name.endswith(".parcial")}
    nb, sv = arvore(NOTEBOOK / "copias"), arvore(SERVIDOR)
    diferentes = sorted(k for k in nb if sv.get(k) != nb[k])
    so_estragado = all(("SERVIDOR:" + Path(k).name) in ESTRAGADOS for k in diferentes)
    r.append(("Servidor tem tudo o que o notebook tem", set(nb) <= set(sv), f"{len(nb)} arquivos"))
    r.append(("Servidor igual ao notebook, byte a byte", so_estragado,
              "a única diferença é o arquivo estragado de propósito (não é consertado sozinho, só avisado)" if diferentes
              else "idênticos"))

    indice = bc.ler_indice(NOTEBOOK / "copias")["arquivos"]
    faltam = [i for i in COMPROVANTES_ENVIADOS if i not in indice]
    errados = [i for i in COMPROVANTES_ENVIADOS if i in indice and hashlib.md5(
        (NOTEBOOK / "copias" / indice[i]["arquivo"]).read_bytes()).hexdigest() != COMPROVANTES_ENVIADOS[i]]
    mais_de_uma = {i for i in DRIVE.baixados if DRIVE.baixados.count(i) > 1}
    r.append(("Comprovantes: todos os enviados estão no backup", not faltam, f"{len(indice)} de {len(COMPROVANTES_ENVIADOS)}"))
    r.append(("Comprovantes: conteúdo idêntico ao original", not errados, "ok" if not errados else f"{len(errados)} diferentes"))
    r.append(("Comprovantes: cada um baixado UMA vez só", not mais_de_uma, f"{len(DRIVE.baixados)} downloads"))
    no_mes = all(Path(v["arquivo"]).parts[1:3] == (Path(v["arquivo"]).name[:4],
                 backup.MESES_PASTA[int(Path(v["arquivo"]).name[5:7]) - 1]) for v in indice.values())
    r.append(("Comprovantes: cada um na pasta do mês do pagamento", no_mes, "ok"))
    from openpyxl import load_workbook
    ws = load_workbook(NOTEBOOK / "copias" / bc.PASTA / bc.LISTA).active
    ligados = sum(1 for l in ws.iter_rows(min_row=2, values_only=True) if "sem parcela ligada" not in (l[1] or ""))
    r.append(("Lista de comprovantes: uma linha por arquivo", ws.max_row - 1 == len(indice),
              f"{ws.max_row - 1} linhas, {ligados} ligadas a uma parcela"))

    # O que o app enxerga numa planilha reconstruída da ÚLTIMA cópia tem que ser o que ele enxerga hoje
    ultima = copias_em(COFRE)[-1][1]["abas"]
    linhas = lambda n: [dict(zip(ultima[n][0], l + [""] * (len(ultima[n][0]) - len(l)))) for l in ultima[n][1:]]
    outro = dados.Local(str(TMP / "reconstruida.local.json"))
    outro._gravar({"cobrancas": linhas("_Cobrancas"), "parcelas": linhas("_Parcelas"),
                   "config": {l[0]: l[1] for l in ultima["_Config"][1:]}})
    igual = outro.carregar() == BANCO.carregar() and outro.get_config() == BANCO.get_config()
    r.append(("Planilha reconstruída da última cópia: o app enxerga o mesmo", igual, f"{len(BANCO.carregar()['cobrancas'])} cobranças"))

    def esperado(quando, texto):
        dia = quando.date()
        return (("servidor" in texto and entre(dia, SERVIDOR_FORA[0] + timedelta(days=2), SERVIDOR_FORA[1] + timedelta(days=1)))
                or ("ler a planilha" in texto and entre(dia, LEITURA_FORA[0] + timedelta(days=2), LEITURA_FORA[1] + timedelta(days=1)))
                or ("Drive" in texto and entre(dia, DRIVE_FORA[0] + timedelta(days=2), DRIVE_FORA[1] + timedelta(days=1)))
                or ("comprovantes" in texto and entre(dia, DRIVE_FORA[0] + timedelta(days=2), DRIVE_FORA[1] + timedelta(days=1)))
                or ("diferentes do original" in texto and dia >= date(2026, 12, 21))
                or "encolheu" in texto)
    falsos = [(q.strftime("%d/%m %H:%M"), l) for q, t in AVISOS for l in t.splitlines()
              if l.strip() and "Abra o Claude" not in l and not esperado(q, l)]
    tem = lambda palavra: any(palavra in t for _, t in AVISOS)
    r.append(("Aviso na tela: servidor 6 dias fora do ar", tem("entregar a cópia ao servidor"), "apareceu a partir do 3º dia"))
    r.append(("Aviso na tela: planilha 4 dias sem abrir", tem("ler a planilha"), "apareceu a partir do 3º dia"))
    r.append(("Aviso na tela: Drive 4 dias fora do ar", tem("gravar a cópia no Drive"), "apareceu a partir do 3º dia"))
    r.append(("Aviso na tela: arquivo do servidor estragado", tem("diferentes do original"), "o backup percebeu e avisou"))
    r.append(("Nenhum aviso falso na tela", not falsos, f"{len(AVISOS)} avisos, todos nas falhas simuladas" if not falsos else str(falsos[:4])))
    return r


def tamanho(raiz: Path) -> tuple[int, float]:
    arqs = [a for a in raiz.rglob("*") if a.is_file()]
    return len(arqs), sum(a.stat().st_size for a in arqs) / 1e6


def main() -> int:
    t0 = time.perf_counter()
    simular()
    resultado = conferir()
    d = BANCO._ler()
    print(f"\nSIMULAÇÃO DE {(FIM - INICIO).days + 1} DIAS ({INICIO:%d/%m/%Y} a {FIM:%d/%m/%Y}), feita em {time.perf_counter() - t0:.0f} s\n")
    print("USO SIMULADO DO APP")
    print(f"  {CONTAGEM['cobrancas']} cobranças novas, {CONTAGEM['pagamentos']} pagamentos, {CONTAGEM['comprovantes']} comprovantes, "
          f"{CONTAGEM['observacoes']} observações, {CONTAGEM['renegociacoes']} renegociações, "
          f"{CONTAGEM['exclusoes']} exclusões, {CONTAGEM['restauracoes']} restaurações")
    print(f"  planilha no fim: {len(d['cobrancas'])} cobranças, {len(d['parcelas'])} parcelas")
    print(f"\nRODADAS: {RODADAS['notebook']} do notebook, {RODADAS['nuvem']} da nuvem "
          f"({len(RODADAS['nuvem_com_erro'])} da nuvem saíram com erro)")
    motivos = {}
    for _, texto in RODADAS["nuvem_com_erro"]:
        chave = "planilha encolheu" if "encolheu" in texto else "Drive fora do ar" if "Drive" in texto else "planilha não abriu"
        motivos[chave] = motivos.get(chave, 0) + 1
    print("  motivos:", ", ".join(f"{n}x {m}" for m, n in motivos.items()) or "nenhum")
    tempos = RODADAS["notebook_tempo"]
    print(f"  tempo de uma rodada do notebook: {1000 * sum(tempos[:10]) / 10:.0f} ms no começo, "
          f"{1000 * sum(tempos[-10:]) / 10:.0f} ms no fim (disco local, sem internet)")
    print("\nO QUE FICOU GUARDADO")
    for nome, raiz in (("notebook", NOTEBOOK / "copias"), ("servidor", SERVIDOR), ("cofre do GitHub", COFRE)):
        n, mb = tamanho(raiz)
        planilha = sum(a.stat().st_size for a in raiz.rglob(f"*{backup.SUFIXO}.*")) / 1e6
        print(f"  {nome:16} {n:5} arquivos, {mb:7.1f} MB (cópias da planilha: {planilha:.1f} MB)")
    nd = [a for a in DRIVE.arqs.values() if backup.SUFIXO in a["name"]]
    print(f"  {'Drive (backup)':16} {len(nd):5} arquivos, {sum(len(a['conteudo']) for a in nd) / 1e6:7.1f} MB")
    print("\nQUANTO TEMPO UMA MUDANÇA ESPEROU ATÉ TER CÓPIA (horas)")
    for nome in ("qualquer",) + DESTINOS:
        e = sorted(ESPERAS[nome])
        if e:
            rotulo = "em algum lugar" if nome == "qualquer" else nome
            print(f"  {rotulo:15} típica {e[len(e) // 2]:5.1f} h   9 em cada 10 até {e[int(len(e) * 0.9)]:5.1f} h   pior caso {e[-1]:6.1f} h")
    print("\nCONFERÊNCIA")
    for nome, ok, detalhe in resultado:
        print(f"  [{'OK' if ok else 'FALHOU'}] {nome}: {detalhe}")
    falhas = [n for n, ok, _ in resultado if not ok]
    print(f"\n{len(resultado) - len(falhas)} de {len(resultado)} conferências passaram." + ("" if not falhas else " FALHARAM: " + "; ".join(falhas)))
    if "--manter" in sys.argv:
        print("Pasta da simulação:", TMP)
    else:
        shutil.rmtree(TMP, ignore_errors=True)
    return 1 if falhas else 0


if __name__ == "__main__":
    sys.exit(main())
