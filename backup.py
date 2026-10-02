#!/usr/bin/env python3
"""
Cópia de segurança da planilha de cobranças. Só LÊ a planilha, nunca escreve nela.

Cada cópia é a planilha INTEIRA (todas as abas), em dois arquivos iguais em conteúdo:
  AAAA-MM-DD HHhMM backup cobrancas.xlsx   para abrir no Excel ou no Google Planilhas
  AAAA-MM-DD HHhMM backup cobrancas.json   o mesmo, exato, para reconstruir a planilha
                                     (ferramentas/restaurar_backup.py)

Regras:
  - cópia nova só quando algo MUDOU desde a última (a impressão digital do conteúdo
    é comparada); sem mudança, só atualiza o ULTIMA_CONFERENCIA.txt;
  - nenhuma cópia antiga é apagada nem reescrita (duas cópias diferentes no mesmo minuto: a
    nova ganha o minuto seguinte no nome);
  - se a planilha ENCOLHEU (cobrança sumiu, aba sumiu, muita parcela sumiu), a cópia
    nova é guardada do mesmo jeito, mas o programa sai com erro para alguém olhar.

Três destinos, independentes entre si:
  --pasta DIR     grava em DIR/AAAA/NN - Mês/ (o cofre do GitHub e a pasta do notebook)
  --espelho DIR   leva para DIR tudo que estiver em --pasta e ainda não estiver lá
                  (a pasta do servidor da empresa, que pode estar fora do ar)
  --drive         grava no Drive da empresa, pasta BACKUP / BACKUP COBRANCAS
                  (a mesma organização do servidor: uma pasta geral, uma subpasta por sistema)

Variáveis de ambiente:
  SPREADSHEET_ID       planilha
  GCP_SERVICE_ACCOUNT  JSON da conta de serviço (senão lê credentials.json)
  DRIVE_OAUTH          JSON {"client_id", "client_secret", "refresh_token"} da conta
                       dona do Drive; sem ele, --drive é pulado com um aviso

Sai com 0 se tudo certo, 1 se alguma cópia falhou ou a planilha encolheu.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import shutil
import sys
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

TIMEZONE = ZoneInfo("America/Sao_Paulo")
# Só leitura: o backup nunca escreve na planilha
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]

ABAS_OBRIGATORIAS = ("_Cobrancas", "_Parcelas")
# O nome começa pela data (ano na frente): na pasta, a ordem por nome já é a ordem do tempo
SUFIXO = " backup cobrancas"
# Uma pasta por mês dentro da pasta do ano (mesmo formato das pastas de comprovante)
MESES_PASTA = ["01 - Janeiro", "02 - Fevereiro", "03 - Março", "04 - Abril", "05 - Maio", "06 - Junho",
               "07 - Julho", "08 - Agosto", "09 - Setembro", "10 - Outubro", "11 - Novembro",
               "12 - Dezembro"]
MARCA = "ULTIMA_CONFERENCIA.txt"
ABA_LEIAME = "LEIA-ME"
# Abas do Excel feitas para gente ler (amarelas); as outras são a cópia exata da planilha
ABAS_LEITURA = ("Resumo", "Parcela por parcela")
COR_LEITURA = "F2DD82"
# Sobe quando o arquivo de backup ganha algo novo: a próxima rodada grava uma cópia no formato
# novo mesmo sem mudança na planilha (2 = o Excel ganhou as abas de leitura)
FORMATO = 2
PASTA_DRIVE_GERAL = "BACKUP"          # pasta geral de backups (outros sistemas ganham a própria subpasta)
PASTA_DRIVE = "BACKUP COBRANCAS"      # mesmo nome da pasta do servidor
# Parcela pode ser removida pelo app (editar cobrança); só acusa sumiço em massa
QUEDA_PARCELAS = 0.20

_MIME_PASTA = "application/vnd.google-apps.folder"
_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_REAIS = "R$ #,##0.00"
_DATA = "DD/MM/YYYY"
_ROTULO_PARCELA = {"paga": "Paga", "atrasada": "Atrasada", "em_breve": "Vence em breve", "aberta": "Em aberto"}
# (começo do texto da situação, cor de fundo)
_COR_SITUACAO = [("Atrasad", "F8D7DA"), ("Vence em breve", "FFE5B4"), ("Quitado", "D4EDDA"),
                 ("Paga", "D4EDDA"), ("Excluída", "E2E3E5")]


def _limpo(valor: str) -> str:
    """Tira espaços e o BOM (marca invisível que o PowerShell põe no começo ao gravar um secret)."""
    return (valor or "").replace("﻿", "").strip()


def agora_sp() -> datetime:
    return datetime.now(tz=TIMEZONE)


# ── Leitura ───────────────────────────────────────────────────────────────────

def normalizar(abas: dict) -> dict:
    """Tudo texto, sem células vazias sobrando no fim da linha nem linhas vazias no fim da aba."""
    saida = {}
    for nome, linhas in abas.items():
        limpas = []
        for linha in linhas or []:
            linha = ["" if c is None else str(c) for c in linha]
            while linha and linha[-1] == "":
                linha.pop()
            limpas.append(linha)
        while limpas and not limpas[-1]:
            limpas.pop()
        saida[str(nome)] = limpas
    return saida


def impressao_digital(abas: dict) -> str:
    """Igual para conteúdo igual: é o que decide se vale gravar uma cópia nova."""
    texto = json.dumps(normalizar(abas), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def ler_planilha(spreadsheet_id: str) -> dict:
    """Todas as abas da planilha, como texto (o app grava tudo como texto)."""
    import gspread
    from google.oauth2.service_account import Credentials

    raw = _limpo(os.environ.get("GCP_SERVICE_ACCOUNT", ""))
    if not raw:
        with open("credentials.json", encoding="utf-8") as f:
            raw = f.read()
    gc = gspread.authorize(Credentials.from_service_account_info(json.loads(raw), scopes=SCOPES))
    sh = gc.open_by_key(spreadsheet_id)
    titulos = [ws.title for ws in sh.worksheets()]
    faixas = ["'" + t.replace("'", "''") + "'" for t in titulos]
    blocos = sh.values_batch_get(faixas).get("valueRanges", [])
    if len(blocos) != len(titulos):
        raise RuntimeError(f"a planilha tem {len(titulos)} abas e a leitura devolveu {len(blocos)}")
    return montar(sh.title, {t: b.get("values", []) for t, b in zip(titulos, blocos)})


def montar(titulo: str, abas: dict, quando: datetime | None = None) -> dict:
    abas = normalizar(abas)
    return {"formato": FORMATO, "planilha": titulo,
            "feito_em": (quando or agora_sp()).isoformat(timespec="seconds"),
            "impressao": impressao_digital(abas), "abas": abas}


def contagem(copia: dict) -> dict:
    """Linhas de dados (sem o cabeçalho) de cada aba."""
    return {nome: max(len(linhas) - 1, 0) for nome, linhas in copia["abas"].items()}


def conferir(copia: dict) -> list[str]:
    """Problemas que fazem a cópia não servir (planilha errada ou quebrada)."""
    problemas = []
    for nome in ABAS_OBRIGATORIAS:
        linhas = copia["abas"].get(nome)
        if linhas is None:
            problemas.append(f"a aba {nome} não existe na planilha")
        elif not linhas or not any(linhas[0]):
            problemas.append(f"a aba {nome} está sem cabeçalho")
    return problemas


def encolheu(antes: dict | None, agora: dict) -> list[str]:
    """O que sumiu em relação à cópia anterior (o app nunca apaga cobrança, só marca)."""
    if not antes:
        return []
    a, b = contagem(antes), contagem(agora)
    avisos = [f"a aba {nome} sumiu da planilha" for nome in a if nome not in b and nome != ABA_LEIAME]
    if b.get("_Cobrancas", 0) < a.get("_Cobrancas", 0):
        avisos.append(f"_Cobrancas tinha {a['_Cobrancas']} linhas e agora tem {b.get('_Cobrancas', 0)}")
    pa, pb = a.get("_Parcelas", 0), b.get("_Parcelas", 0)
    if pa and pb < pa * (1 - QUEDA_PARCELAS):
        avisos.append(f"_Parcelas tinha {pa} linhas e agora tem {pb}")
    return avisos


# ── Arquivos ──────────────────────────────────────────────────────────────────

def nome_arquivo(copia: dict) -> str:
    quando = datetime.fromisoformat(copia["feito_em"])
    return quando.strftime("%Y-%m-%d %Hh%M") + SUFIXO


def no_minuto_livre(copia: dict, ocupado) -> dict:
    """
    Se já existe cópia com este nome (a planilha mudou duas vezes no mesmo minuto), a nova ganha
    o minuto seguinte. Assim nenhuma cópia antiga pode ser sobrescrita.
    """
    while ocupado(nome_arquivo(copia)):
        quando = datetime.fromisoformat(copia["feito_em"]) + timedelta(minutes=1)
        copia = {**copia, "feito_em": quando.isoformat(timespec="seconds")}
    return copia


def pastas_da_copia(nome: str) -> tuple[str, str]:
    """(ano, mês) onde a cópia mora, tirados do próprio nome: "2026", "10 - Outubro"."""
    return nome[:4], MESES_PASTA[int(nome[5:7]) - 1]


def gerar_json(copia: dict) -> bytes:
    return json.dumps(copia, ensure_ascii=False, indent=1).encode("utf-8")


def _escrever(cel, valor):
    """Texto vai sempre como TEXTO (nada vira fórmula); número e data vão como estão."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if isinstance(valor, str):
        if valor == "":
            return
        cel.value = ILLEGAL_CHARACTERS_RE.sub("", valor)
        cel.data_type = "s"
    elif valor is not None:
        cel.value = valor


def _nome_livre(nome: str, usados: set) -> str:
    while nome.lower() in usados:
        nome += " (leitura)"
    usados.add(nome.lower())
    return nome[:31]


def _tabela(ws, linha0: int, colunas: list, linhas: list):
    """colunas = [(título, largura, formato)]; a coluna "Situação" ganha cor conforme o texto."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cabecalho = PatternFill("solid", fgColor="1F3A5F")
    for j, (titulo, largura, _) in enumerate(colunas, start=1):
        cel = ws.cell(row=linha0, column=j, value=titulo)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = cabecalho
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = largura
    for i, linha in enumerate(linhas, start=linha0 + 1):
        for j, valor in enumerate(linha, start=1):
            cel = ws.cell(row=i, column=j)
            _escrever(cel, valor)
            formato = colunas[j - 1][2]
            if formato and not isinstance(valor, str):
                cel.number_format = formato
            if colunas[j - 1][0] == "Situação" and isinstance(valor, str):
                cor = next((c for chave, c in _COR_SITUACAO if valor.startswith(chave)), None)
                if cor:
                    cel.fill = PatternFill("solid", fgColor=cor)
    ws.freeze_panes = ws.cell(row=linha0 + 1, column=2)
    if linhas:
        ws.auto_filter.ref = f"A{linha0}:{get_column_letter(len(colunas))}{linha0 + len(linhas)}"


def _abas_de_leitura(wb, copia: dict, usados: set):
    """
    Duas abas para GENTE ler, com as mesmas contas do app (pago, falta, atrasado, situação).
    As outras abas são a cópia exata, e é só delas que a planilha é reconstruída.
    """
    import nucleo
    from openpyxl.styles import Font

    quando = datetime.fromisoformat(copia["feito_em"])
    abas = copia["abas"]
    cobs = nucleo.montar_cobrancas(nucleo.linhas_de_valores(abas.get("_Cobrancas", [])),
                                   nucleo.linhas_de_valores(abas.get("_Parcelas", [])),
                                   quando.date(), incluir_excluidas=True)
    ativas = nucleo.ordenar_para_cobrar([c for c in cobs if not c["excluida"]])
    excluidas = [c for c in cobs if c["excluida"]]
    totais = nucleo.resumo(ativas)

    def situacao(c):
        if c["excluida"]:
            return "Excluída"
        if c["situacao"] == nucleo.CLI_ATRASADO:
            dias = c["dias_atraso"]
            return f"Atrasado há {dias} dia{'s' if dias != 1 else ''}"
        return nucleo.ROTULO_CLI[c["situacao"]]

    resumo = wb.create_sheet(title=_nome_livre(ABAS_LEITURA[0], usados))
    resumo.sheet_properties.tabColor = COR_LEITURA
    resumo["A1"].value = "Cobranças Vithall"
    resumo["A1"].font = Font(bold=True, size=14)
    resumo["A2"].value = f"Situação em {quando.strftime('%d/%m/%Y às %H:%M')} (horário de Brasília)"
    topo = [("Cobranças ativas", len(ativas), "0"), ("Já recebido", totais["recebido"], _REAIS),
            ("Falta receber", totais["a_receber"], _REAIS), ("Atrasado", totais["atrasado"], _REAIS),
            ("Clientes atrasados", totais["n_atrasados"], "0")]
    for i, (rotulo, valor, formato) in enumerate(topo, start=4):
        resumo.cell(row=i, column=1, value=rotulo).font = Font(bold=True)
        resumo.cell(row=i, column=2, value=valor).number_format = formato
    colunas = [("Cliente", 34, None), ("Turma", 9, None), ("Cidade", 16, None), ("Treinamento", 13, None),
               ("Forma", 9, None), ("Valor total", 14, _REAIS), ("Já pago", 14, _REAIS),
               ("Falta receber", 14, _REAIS), ("Parcelas pagas", 10, None), ("Próxima cobrança", 13, _DATA),
               ("Valor da próxima", 14, _REAIS), ("Situação", 22, None), ("Observações", 40, None)]
    linhas = []
    for c in ativas + excluidas:
        ps, prox = c["parcelas"], c["proxima"]
        total = c["valor_total"] if c["valor_total"] is not None else c["pago"] + c["falta"]
        linhas.append([c["cliente"], c["turma"], c["cidade"], c["treinamento"], c["tipo"], total, c["pago"],
                       c["falta"], f"{sum(1 for p in ps if p['paga'])} de {len(ps)}",
                       prox["vencimento"] if prox else None, prox["valor"] if prox else None,
                       situacao(c), c["observacoes"]])
    _tabela(resumo, 10, colunas, linhas)

    parcelas = wb.create_sheet(title=_nome_livre(ABAS_LEITURA[1], usados))
    parcelas.sheet_properties.tabColor = COR_LEITURA
    colunas = [("Cliente", 34, None), ("Turma", 9, None), ("Cidade", 16, None), ("Forma", 9, None),
               ("Parcela", 18, None), ("Vencimento", 13, _DATA), ("Valor", 14, _REAIS), ("Situação", 16, None),
               ("Pago em", 13, _DATA), ("Quem recebeu", 16, None), ("Observação", 36, None),
               ("Comprovante", 40, None)]
    linhas = [[c["cliente"], c["turma"], c["cidade"], c["tipo"], p["rotulo"], p["vencimento"], p["valor"],
               _ROTULO_PARCELA[p["situacao"]], p["pago_em"], p["marcado_por"], p["observacao"], p["comprovante"]]
              for c in ativas for p in c["parcelas"]]
    _tabela(parcelas, 1, colunas, linhas)


def gerar_xlsx(copia: dict) -> bytes:
    """
    Abas amarelas (Resumo, Parcela por parcela, LEIA-ME): para gente ler.
    Depois, uma aba por aba da planilha, toda célula como TEXTO (nada vira número, data ou
    fórmula): é a cópia exata, de onde a planilha é reconstruída.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    usados = {nome[:31].lower() for nome in copia["abas"]} | {ABA_LEIAME.lower()}
    sem_resumo = ""
    try:
        _abas_de_leitura(wb, copia, usados)
    except Exception as e:   # o resumo é um extra: nunca pode impedir a cópia exata
        for ws in list(wb.worksheets):
            wb.remove(ws)
        sem_resumo = f"As abas de leitura não puderam ser montadas nesta cópia ({type(e).__name__}: {e})."

    leia = wb.create_sheet(title=ABA_LEIAME)
    leia.sheet_properties.tabColor = COR_LEITURA
    quando = datetime.fromisoformat(copia["feito_em"])
    textos = [f"Cópia de segurança da planilha \"{copia['planilha']}\"",
              f"Feita em {quando.strftime('%d/%m/%Y às %H:%M')} (horário de Brasília)", "",
              "ABAS AMARELAS (Resumo e Parcela por parcela): para ler. Mostram cada cliente, quanto pagou,",
              "quanto falta e a situação, com as mesmas contas do app. Pode filtrar e ordenar à vontade.", "",
              "ABAS QUE COMEÇAM COM _ (e as demais sem cor): cópia exata da planilha do app. Não mexa nelas:",
              "é a partir delas que a planilha é reconstruída se a original for perdida.", ""]
    if sem_resumo:
        textos += [sem_resumo, ""]
    textos += ["Linhas em cada aba da cópia exata:"]
    textos += [f"  {nome}: {n}" for nome, n in contagem(copia).items()]
    textos += ["", "Para recuperar: abra este arquivo no Google Planilhas (Arquivo > Importar) ou",
               "use o arquivo .json de mesmo nome com ferramentas/restaurar_backup.py.",
               "Impressão digital do conteúdo: " + copia["impressao"]]
    for i, texto in enumerate(textos, start=1):
        leia.cell(row=i, column=1).value = texto
    leia.column_dimensions["A"].width = 100

    for nome, linhas in copia["abas"].items():
        ws = wb.create_sheet(title=nome[:31])
        for i, linha in enumerate(linhas, start=1):
            for j, valor in enumerate(linha, start=1):
                if valor == "":
                    continue
                cel = ws.cell(row=i, column=j)
                _escrever(cel, valor)
                cel.number_format = "@"
                if i == 1:
                    cel.font = Font(bold=True)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ler_xlsx(conteudo: bytes) -> dict:
    """Só a cópia exata de um .xlsx de backup, de volta ao formato da planilha (pula as abas amarelas)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo))
    abas = {ws.title: [list(linha) for linha in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets if ws.sheet_properties.tabColor is None}
    return normalizar(abas)


def ultima_copia(pasta: Path) -> dict | None:
    """A cópia mais recente guardada na pasta, em qualquer subpasta (o nome do arquivo já ordena por data)."""
    for arq in sorted(Path(pasta).rglob(f"*{SUFIXO}.json"), key=lambda p: p.name, reverse=True):
        try:
            return json.loads(arq.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue   # arquivo cortado no meio: vale o anterior
    return None


def _gravar(caminho: Path, conteudo: bytes):
    """Grava inteiro ou não grava: escreve ao lado e troca o nome no fim."""
    caminho.parent.mkdir(parents=True, exist_ok=True)
    temporario = caminho.with_name(caminho.name + ".parcial")
    temporario.write_bytes(conteudo)
    os.replace(temporario, caminho)


def texto_marca(copia: dict, arquivo: str) -> str:
    quando = datetime.fromisoformat(copia["feito_em"])
    partes = ", ".join(f"{n} em {nome}" for nome, n in contagem(copia).items())
    return (f"Planilha conferida em {quando.strftime('%d/%m/%Y às %H:%M')} (horário de Brasília).\n"
            f"Linhas: {partes}.\n"
            f"Cópia mais recente: {arquivo}\n"
            "Cópia nova só é gravada quando algo muda na planilha; este arquivo é atualizado toda vez.\n")


def salvar_pasta(copia: dict, pasta: Path, marca: bool = True) -> tuple[str, bool, list[str]]:
    """Guarda a cópia em pasta/AAAA/NN - Mês/. Devolve (nome do arquivo que vale, gravou agora?, o que sumiu)."""
    pasta = Path(pasta)
    anterior = ultima_copia(pasta)
    sumiu = encolheu(anterior, copia)
    if (anterior and anterior.get("impressao") == copia["impressao"]
            and anterior.get("formato") == copia["formato"]):
        nome, gravou = nome_arquivo(anterior), False
    else:
        copia = no_minuto_livre(copia, lambda n: any(
            Path(pasta, *pastas_da_copia(n), n + ext).exists() for ext in (".xlsx", ".json")))
        nome, gravou = nome_arquivo(copia), True
        destino = Path(pasta, *pastas_da_copia(nome))
        _gravar(destino / (nome + ".xlsx"), gerar_xlsx(copia))
        _gravar(destino / (nome + ".json"), gerar_json(copia))   # o .json por último: é ele que marca "cópia completa"
    if marca:
        _gravar(pasta / MARCA, texto_marca(copia, nome + ".xlsx").encode("utf-8"))
    return nome, gravou, sumiu


def espelhar(origem: Path, destino: Path) -> int:
    """Leva para destino o que existe em origem e falta lá. Nunca apaga; cópia datada nunca é reescrita."""
    origem, destino = Path(origem), Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    levados = 0
    for arq in sorted(origem.rglob("*")):
        if not arq.is_file() or arq.name.endswith(".parcial"):
            continue
        alvo = destino / arq.relative_to(origem)
        fixo = SUFIXO in arq.name   # cópia datada nunca muda; os outros (marca, leia-me) mudam
        if alvo.exists() and (fixo or alvo.read_bytes() == arq.read_bytes()):
            continue
        alvo.parent.mkdir(parents=True, exist_ok=True)
        parcial = alvo.with_name(alvo.name + ".parcial")
        shutil.copyfile(arq, parcial)
        os.replace(parcial, alvo)
        levados += 1
    return levados


# ── Drive da empresa ──────────────────────────────────────────────────────────

def credencial_drive() -> dict | None:
    raw = _limpo(os.environ.get("DRIVE_OAUTH", ""))
    if not raw:
        return None
    dados = json.loads(raw)
    return {k: dados[k] for k in ("client_id", "client_secret", "refresh_token")}


def servico_drive(cred: dict):
    """API do Drive como o DONO da conta; só enxerga o que este mesmo acesso criou (drive.file)."""
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build
    creds = Credentials(token=None, refresh_token=cred["refresh_token"], client_id=cred["client_id"],
                        client_secret=cred["client_secret"],
                        token_uri="https://oauth2.googleapis.com/token",
                        scopes=["https://www.googleapis.com/auth/drive.file"])
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def _achar_pasta(service, nome: str, pai: str | None) -> str | None:
    """Sem pai, procura a pasta ONDE ELA ESTIVER no Drive (a mais antiga, se houver duas)."""
    seguro = nome.replace("\\", "\\\\").replace("'", "\\'")
    onde = f" and '{pai}' in parents" if pai else ""
    achados = service.files().list(
        q=f"name = '{seguro}'{onde} and mimeType = '{_MIME_PASTA}' and trashed = false",
        orderBy="createdTime", fields="files(id)", pageSize=1).execute().get("files", [])
    return achados[0]["id"] if achados else None


def _pasta_drive(service, nome: str, pai: str | None) -> str:
    achada = _achar_pasta(service, nome, pai)
    if achada:
        return achada
    corpo = {"name": nome, "mimeType": _MIME_PASTA}
    if pai:
        corpo["parents"] = [pai]
    return service.files().create(body=corpo, fields="id").execute()["id"]


def pasta_backup_drive(service) -> str:
    """
    A pasta das cópias. Vale onde ela estiver (o dono pode arrastá-la de lugar sem quebrar nada);
    se não existir em lugar nenhum, nasce em BACKUP / BACKUP COBRANCAS.
    """
    return (_achar_pasta(service, PASTA_DRIVE, None)
            or _pasta_drive(service, PASTA_DRIVE, _pasta_drive(service, PASTA_DRIVE_GERAL, None)))


def enviar_drive(copia: dict, service) -> tuple[str, bool]:
    """Mesma regra da pasta: cópia nova só se mudou; a marca é atualizada sempre."""
    from googleapiclient.http import MediaIoBaseUpload

    raiz = pasta_backup_drive(service)
    recentes = service.files().list(
        q="appProperties has { key='tipo' and value='backup-cobrancas' } and trashed = false",
        orderBy="name desc", pageSize=1, fields="files(name, appProperties)").execute().get("files", [])
    props = recentes[0].get("appProperties", {}) if recentes else {}
    if props.get("impressao") == copia["impressao"] and props.get("formato") == str(copia["formato"]):
        nome, gravou = recentes[0]["name"].rsplit(".", 1)[0], False
    else:
        def ocupado(n):
            return bool(service.files().list(
                q=f"(name = '{n}.xlsx' or name = '{n}.json') and trashed = false",
                fields="files(id)", pageSize=1).execute().get("files"))
        copia = no_minuto_livre(copia, ocupado)
        nome, gravou = nome_arquivo(copia), True
        ano, mes = pastas_da_copia(nome)
        destino = _pasta_drive(service, mes, _pasta_drive(service, ano, raiz))
        # Só o .json (enviado por último) leva a etiqueta: se o envio cair no meio, a próxima rodada refaz
        etiqueta = {"tipo": "backup-cobrancas", "impressao": copia["impressao"],
                    "formato": str(copia["formato"])}
        for ext, conteudo, mime, marcas in ((".xlsx", gerar_xlsx(copia), _MIME_XLSX, {}),
                                           (".json", gerar_json(copia), "application/json", etiqueta)):
            service.files().create(
                body={"name": nome + ext, "parents": [destino], "appProperties": marcas},
                media_body=MediaIoBaseUpload(io.BytesIO(conteudo), mimetype=mime, resumable=False),
                fields="id").execute()
    marca = MediaIoBaseUpload(io.BytesIO(texto_marca(copia, nome + ".xlsx").encode("utf-8")),
                              mimetype="text/plain", resumable=False)
    existentes = service.files().list(
        q=f"name = '{MARCA}' and '{raiz}' in parents and trashed = false",
        fields="files(id)", pageSize=1).execute().get("files", [])
    if existentes:
        service.files().update(fileId=existentes[0]["id"], media_body=marca).execute()
    else:
        service.files().create(body={"name": MARCA, "parents": [raiz]}, media_body=marca).execute()
    return nome, gravou


# ── Execução ──────────────────────────────────────────────────────────────────

def executar(copia: dict, pasta: Path | None = None, espelho: Path | None = None,
             drive: bool = False, saida=print) -> list[str]:
    """Roda os destinos pedidos; uma falha num destino não impede os outros. Devolve os erros."""
    erros = list(conferir(copia))
    if erros:
        return erros   # planilha errada ou quebrada: não vale como cópia
    saida("Planilha lida: " + ", ".join(f"{n} em {nome}" for nome, n in contagem(copia).items()))
    if pasta:
        try:
            nome, gravou, sumiu = salvar_pasta(copia, pasta)
            saida(f"Pasta: {'cópia nova ' + nome if gravou else 'sem mudança desde ' + nome}")
            erros += [f"ATENÇÃO, a planilha encolheu: {s}" for s in sumiu]
        except Exception as e:
            erros.append(f"pasta {pasta}: {type(e).__name__}: {e}")
    if espelho:
        if not pasta:
            erros.append("--espelho precisa de --pasta")
        else:
            try:
                saida(f"Espelho: {espelhar(pasta, espelho)} arquivo(s) levado(s) para {espelho}")
            except Exception as e:
                erros.append(f"espelho {espelho}: {type(e).__name__}: {e}")
    if drive:
        try:
            cred = credencial_drive()
            if cred is None:
                saida("Drive: DRIVE_OAUTH não configurado, pulando")
            else:
                nome, gravou = enviar_drive(copia, servico_drive(cred))
                saida(f"Drive: {'cópia nova ' + nome if gravou else 'sem mudança desde ' + nome}")
        except Exception as e:
            erros.append(f"Drive: {type(e).__name__}: {e}")
    return erros


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Cópia de segurança da planilha de cobranças")
    ap.add_argument("--pasta", type=Path)
    ap.add_argument("--espelho", type=Path)
    ap.add_argument("--drive", action="store_true")
    args = ap.parse_args(argv)
    spreadsheet_id = _limpo(os.environ.get("SPREADSHEET_ID", ""))
    if not spreadsheet_id:
        print("ERRO: SPREADSHEET_ID não definido", file=sys.stderr)
        return 1
    try:
        copia = ler_planilha(spreadsheet_id)
    except Exception as e:
        print(f"ERRO ao ler a planilha: {type(e).__name__}: {e}", file=sys.stderr)
        return 1
    erros = executar(copia, args.pasta, args.espelho, args.drive)
    for e in erros:
        print("ERRO: " + e, file=sys.stderr)
    return 1 if erros else 0


if __name__ == "__main__":
    sys.exit(main())
