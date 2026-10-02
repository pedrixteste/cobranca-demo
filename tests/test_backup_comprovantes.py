import hashlib
import io
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import backup as b  # noqa: E402
import backup_comprovantes as bc  # noqa: E402

AGORA = datetime(2026, 10, 5, 10, 30, tzinfo=b.TIMEZONE)


def arq(fid, nome, conteudo=b"foto", criado="2026-10-02T13:00:00.000Z", pasta="LORAP/2026/10 - Outubro"):
    return {"id": fid, "name": nome, "createdTime": criado, "pasta": pasta,
            "md5Checksum": hashlib.md5(conteudo).hexdigest(), "webViewLink": f"https://drive.google.com/file/d/{fid}/view",
            "_conteudo": conteudo}


def baixador(arquivos, contador=None):
    por_id = {a["id"]: a["_conteudo"] for a in arquivos}

    def baixar(fid):
        if contador is not None:
            contador.append(fid)
        return por_id[fid]
    return baixar


class Nome(unittest.TestCase):
    def test_data_do_pagamento_vai_para_a_frente(self):
        self.assertEqual(bc.nome_no_backup("Maria Souza - L345 - Parcela 2 de 6 - 02-10-2026.jpg", ""),
                         (date(2026, 10, 2), "2026-10-02 Maria Souza - L345 - Parcela 2 de 6.jpg"))
        self.assertEqual(bc.nome_no_backup("João - V12 - Cartão 1 de 4 - 15-01-2027.PDF", "")[1],
                         "2027-01-15 João - V12 - Cartão 1 de 4.pdf")

    def test_sem_data_no_nome_vale_o_dia_em_que_subiu_em_brasilia(self):
        # 01:30 UTC de 03/10 ainda é 02/10 em Brasília
        self.assertEqual(bc.nome_no_backup("foto solta.png", "2026-10-03T01:30:00.000Z"),
                         (date(2026, 10, 2), "2026-10-02 foto solta.png"))
        self.assertEqual(bc.nome_no_backup("Data impossível - 31-02-2026.jpg", "2026-10-03T15:00:00.000Z")[0],
                         date(2026, 10, 3))

    def test_pasta_do_mes_do_pagamento_e_nome_repetido_ganha_numero(self):
        nome = "Maria - L345 - Entrada - 30-11-2026.jpg"
        primeiro = bc.caminho_no_backup(nome, "", set())
        self.assertEqual(primeiro, "COMPROVANTES/2026/11 - Novembro/2026-11-30 Maria - L345 - Entrada.jpg")
        self.assertEqual(bc.caminho_no_backup(nome, "", {primeiro}),
                         "COMPROVANTES/2026/11 - Novembro/2026-11-30 Maria - L345 - Entrada (2).jpg")


class Copia(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name) / "copias"
        self.addCleanup(self._tmp.cleanup)

    def arquivos(self):
        return sorted(p.relative_to(self.pasta).as_posix() for p in self.pasta.rglob("*") if p.is_file())

    def test_cada_comprovante_e_baixado_uma_vez_so(self):
        drive = [arq("id1", "Maria - L345 - Entrada - 02-10-2026.jpg", b"um"),
                 arq("id2", "João - V12 - Cartão 1 de 4 - 05-11-2026.pdf", b"dois")]
        baixados = []
        self.assertEqual(bc.copiar_novos(drive, baixador(drive, baixados), self.pasta), (2, []))
        self.assertEqual(bc.copiar_novos(drive, baixador(drive, baixados), self.pasta), (0, []))
        self.assertEqual(sorted(baixados), ["id1", "id2"])
        self.assertEqual(self.arquivos(), ["COMPROVANTES/2026/10 - Outubro/2026-10-02 Maria - L345 - Entrada.jpg",
                                           "COMPROVANTES/2026/11 - Novembro/2026-11-05 João - V12 - Cartão 1 de 4.pdf",
                                           "COMPROVANTES/indice.json"])
        drive.append(arq("id3", "Ana - P7 - Parcela 1 de 2 - 06-11-2026.jpg", b"tres"))
        self.assertEqual(bc.copiar_novos(drive, baixador(drive, baixados), self.pasta), (1, []))
        self.assertEqual(baixados.count("id1"), 1)

    def test_sumiu_do_drive_continua_no_backup(self):
        drive = [arq("id1", "Maria - L345 - Entrada - 02-10-2026.jpg", b"um")]
        bc.copiar_novos(drive, baixador(drive), self.pasta)
        self.assertEqual(bc.copiar_novos([], baixador([]), self.pasta), (0, []))
        self.assertTrue((self.pasta / "COMPROVANTES/2026/10 - Outubro/2026-10-02 Maria - L345 - Entrada.jpg").exists())

    def test_dois_arquivos_com_o_mesmo_nome_ficam_os_dois(self):
        drive = [arq("id1", "Maria - L345 - Entrada - 02-10-2026.jpg", b"um", criado="2026-10-02T13:00:00.000Z"),
                 arq("id2", "Maria - L345 - Entrada - 02-10-2026.jpg", b"outro", criado="2026-10-02T14:00:00.000Z")]
        self.assertEqual(bc.copiar_novos(drive, baixador(drive), self.pasta)[0], 2)
        pasta = self.pasta / "COMPROVANTES/2026/10 - Outubro"
        self.assertEqual((pasta / "2026-10-02 Maria - L345 - Entrada.jpg").read_bytes(), b"um")
        self.assertEqual((pasta / "2026-10-02 Maria - L345 - Entrada (2).jpg").read_bytes(), b"outro")

    def test_arquivo_corrompido_no_caminho_nao_e_guardado_e_os_outros_seguem(self):
        ruim = arq("id1", "Maria - L345 - Entrada - 02-10-2026.jpg", b"um")
        ruim["md5Checksum"] = "0" * 32
        bom = arq("id2", "João - V12 - Entrada - 03-10-2026.jpg", b"dois")
        novos, erros = bc.copiar_novos([ruim, bom], baixador([ruim, bom]), self.pasta)
        self.assertEqual((novos, len(erros)), (1, 1))
        self.assertNotIn("id1", bc.ler_indice(self.pasta)["arquivos"])
        # consertado no Drive, entra na rodada seguinte
        ruim["md5Checksum"] = hashlib.md5(b"um").hexdigest()
        self.assertEqual(bc.copiar_novos([ruim, bom], baixador([ruim, bom]), self.pasta), (1, []))

    def test_espelho_leva_comprovante_uma_vez_e_atualiza_so_a_lista(self):
        servidor = Path(self._tmp.name) / "servidor"
        drive = [arq("id1", "Maria - L345 - Entrada - 02-10-2026.jpg", b"um")]
        bc.copiar_novos(drive, baixador(drive), self.pasta)
        mudam = (b.MARCA, "LEIA-ME.txt", bc.INDICE, bc.LISTA)
        self.assertEqual(b.espelhar(self.pasta, servidor, mudam=mudam), 2)   # a foto e o índice
        foto = servidor / "COMPROVANTES/2026/10 - Outubro/2026-10-02 Maria - L345 - Entrada.jpg"
        foto.write_bytes(b"no servidor")
        drive.append(arq("id2", "João - V12 - Entrada - 03-10-2026.jpg", b"dois"))
        bc.copiar_novos(drive, baixador(drive), self.pasta)
        self.assertEqual(b.espelhar(self.pasta, servidor, mudam=mudam), 2)   # a foto nova e o índice atualizado
        self.assertEqual(foto.read_bytes(), b"no servidor")                  # a antiga não é reescrita


class Lista(unittest.TestCase):
    def test_lista_liga_o_arquivo_a_parcela_da_planilha(self):
        from openpyxl import load_workbook
        link = "https://drive.google.com/file/d/id1/view?usp=drivesdk"
        copia = b.montar("T", {
            "_Cobrancas": [["ID", "Tipo", "Cliente", "Turma", "Valor Total", "Situação"],
                           ["a1", "Pix", "Maria Souza", "L345", "2000,00", "Ativa"]],
            "_Parcelas": [["ID Cobrança", "Nº", "Vencimento", "Valor", "Status", "Pago em", "Comprovante", "Marcado por"],
                          ["a1", "0", "01/10/2026", "1000,00", "Paga", "02/10/2026", link, "Gabi"],
                          ["a1", "1", "01/11/2026", "1000,00", "", "", "", ""]],
        }, AGORA)
        indice = {"planilha": "", "arquivos": {
            "id1": {"arquivo": "COMPROVANTES/2026/10 - Outubro/2026-10-02 Maria Souza - L345 - Entrada.jpg",
                    "nome_drive": "Maria Souza - L345 - Entrada - 02-10-2026.jpg", "enviado_em": "2026-10-02T13:00:00.000Z",
                    "link": link},
            "id9": {"arquivo": "COMPROVANTES/2026/10 - Outubro/2026-10-03 Teste.jpg",
                    "nome_drive": "Teste - 03-10-2026.jpg", "enviado_em": "2026-10-03T13:00:00.000Z", "link": ""}}}
        ws = load_workbook(io.BytesIO(bc.gerar_lista(indice, copia))).active
        linhas = [[c.value for c in linha] for linha in ws.iter_rows(min_row=2)]
        self.assertEqual(len(linhas), 2)
        solta, maria = linhas   # mais recente primeiro
        self.assertEqual(maria[1:6], ["Maria Souza", "L345", "Entrada", 1000, "Gabi"])
        self.assertEqual(maria[0].date(), date(2026, 10, 2))
        self.assertEqual(maria[6], "2026\\10 - Outubro\\2026-10-02 Maria Souza - L345 - Entrada.jpg")
        self.assertIn("sem parcela ligada", solta[1])


if __name__ == "__main__":
    unittest.main()
