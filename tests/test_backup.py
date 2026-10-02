import io
import json
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import backup as b  # noqa: E402

DIA1 = datetime(2026, 10, 1, 7, 40, tzinfo=b.TIMEZONE)
DIA2 = datetime(2026, 10, 2, 13, 0, tzinfo=b.TIMEZONE)


def abas(cobrancas=2, parcelas=6, extra=None):
    d = {
        "_Cobrancas": [["ID", "Cliente", "Valor Total"]] +
                      [[f"c{i}", f"Cliente {i}", "1500,00"] for i in range(cobrancas)],
        "_Parcelas": [["ID Cobrança", "Nº", "Vencimento", "Valor"]] +
                     [[f"c{i % max(cobrancas, 1)}", str(i), "05/01/2027", "0250,50"] for i in range(parcelas)],
        "_Config": [["Chave", "Valor"], ["admin", "Pedro"]],
    }
    d.update(extra or {})
    return d


class Impressao(unittest.TestCase):
    def test_vazio_sobrando_nao_muda_a_impressao(self):
        a = abas()
        com_sobra = {k: [l + ["", ""] for l in v] + [[], [""]] for k, v in a.items()}
        self.assertEqual(b.impressao_digital(a), b.impressao_digital(com_sobra))

    def test_qualquer_celula_diferente_muda(self):
        a, outro = abas(), abas()
        outro["_Parcelas"][3][3] = "0250,51"
        self.assertNotEqual(b.impressao_digital(a), b.impressao_digital(outro))

    def test_celula_vazia_no_meio_conta(self):
        self.assertNotEqual(b.impressao_digital({"x": [["a", "", "b"]]}),
                            b.impressao_digital({"x": [["a", "b"]]}))


class Excel(unittest.TestCase):
    def test_volta_identico_inclusive_texto_perigoso(self):
        perigosos = [["Obs"], ["=1+1"], ["0012"], ["1500,00"], ["01/02/2027"], ["+5599999"],
                     ["  espaço na ponta "], ["linha\nquebrada"], ["", "depois de vazio"], ["TRUE"]]
        copia = b.montar("Teste", abas(extra={"_Obs": perigosos}), DIA1)
        self.assertEqual(b.ler_xlsx(b.gerar_xlsx(copia)), copia["abas"])

    def test_abas_de_leitura_nao_entram_na_volta(self):
        copia = b.montar("Teste", abas(), DIA1)
        self.assertEqual(set(b.ler_xlsx(b.gerar_xlsx(copia))), set(copia["abas"]))

    def test_abas_de_leitura_vem_primeiro_e_batem_com_o_app(self):
        from openpyxl import load_workbook
        dados = {
            "_Cobrancas": [["ID", "Tipo", "Cliente", "Turma", "Valor Total", "Situação", "Cidade", "Observações"],
                           ["a1", "Pix", "Maria Souza", "L345", "3000,00", "Ativa", "Lajeado", "=perigo"],
                           ["b2", "Cartão", "João Lima", "V12", "1000,00", "Ativa", "", ""],
                           ["c3", "Pix", "Excluído da Silva", "L345", "500,00", "Excluída", "", ""]],
            "_Parcelas": [["ID Cobrança", "Nº", "Vencimento", "Valor", "Status", "Pago em", "Marcado por"],
                          ["a1", "0", "01/09/2026", "1000,00", "Paga", "01/09/2026", "Gabi"],
                          ["a1", "1", "20/09/2026", "1000,00", "", "", ""],
                          ["a1", "2", "20/10/2026", "1000,00", "", "", ""],
                          ["b2", "1", "15/10/2026", "1000,00", "Paga", "15/09/2026", "Ana"],
                          ["c3", "1", "01/01/2026", "500,00", "", "", ""]],
        }
        wb = load_workbook(io.BytesIO(b.gerar_xlsx(b.montar("T", dados, DIA1))))
        self.assertEqual(wb.sheetnames[:3], ["Resumo", "Parcela por parcela", "LEIA-ME"])
        r = wb["Resumo"]
        topo = {r.cell(row=i, column=1).value: r.cell(row=i, column=2).value for i in range(4, 9)}
        self.assertEqual(topo, {"Cobranças ativas": 2, "Já recebido": 2000, "Falta receber": 2000,
                                "Atrasado": 1000, "Clientes atrasados": 1})
        linhas = [[c.value for c in linha] for linha in r.iter_rows(min_row=11, max_row=13)]
        self.assertEqual([l[0] for l in linhas], ["Maria Souza", "João Lima", "Excluído da Silva"])
        maria = linhas[0]
        self.assertEqual(maria[5:9], [3000, 1000, 2000, "1 de 3"])
        self.assertEqual(maria[9].date(), DIA1.date().replace(day=20))   # próxima: 20/10/2026
        self.assertEqual((maria[11], maria[12]), ("Atrasado há 11 dias", "=perigo"))
        self.assertEqual(r.cell(row=11, column=13).data_type, "s")       # observação não vira fórmula
        self.assertEqual((linhas[1][11], linhas[2][11]), ("Quitado", "Excluída"))
        p = wb["Parcela por parcela"]
        self.assertEqual(p.max_row, 5)   # cabeçalho + 4 parcelas das cobranças ativas
        self.assertEqual([c.value for c in p[2]][3:8] + [p.cell(row=2, column=10).value],
                         ["Pix", "Entrada", p.cell(row=2, column=6).value, 1000, "Paga", "Gabi"])
        self.assertEqual(p.cell(row=3, column=8).value, "Atrasada")

    def test_resumo_quebrado_nao_impede_a_copia_exata(self):
        import nucleo
        original = nucleo.montar_cobrancas
        nucleo.montar_cobrancas = lambda *a, **k: 1 / 0
        try:
            copia = b.montar("Teste", abas(), DIA1)
            self.assertEqual(b.ler_xlsx(b.gerar_xlsx(copia)), copia["abas"])
        finally:
            nucleo.montar_cobrancas = original

    def test_aba_da_planilha_chamada_resumo_nao_se_perde(self):
        copia = b.montar("Teste", abas(extra={"Resumo": [["meu resumo"]]}), DIA1)
        self.assertEqual(b.ler_xlsx(b.gerar_xlsx(copia)), copia["abas"])


class Conferencia(unittest.TestCase):
    def test_planilha_sem_as_abas_do_app_nao_vale(self):
        self.assertTrue(b.conferir(b.montar("Outra", {"Página1": [["a"]]}, DIA1)))
        self.assertTrue(b.conferir(b.montar("Vazia", {"_Cobrancas": [], "_Parcelas": [["ID"]]}, DIA1)))
        self.assertEqual(b.conferir(b.montar("Certa", abas(0, 0), DIA1)), [])

    def test_encolheu(self):
        antes = b.montar("T", abas(5, 36), DIA1)
        self.assertEqual(b.encolheu(None, antes), [])
        self.assertEqual(b.encolheu(antes, b.montar("T", abas(6, 40), DIA2)), [])
        self.assertEqual(b.encolheu(antes, b.montar("T", abas(5, 34), DIA2)), [])   # tirou 2 parcelas: normal
        self.assertEqual(len(b.encolheu(antes, b.montar("T", abas(4, 36), DIA2))), 1)
        self.assertEqual(len(b.encolheu(antes, b.montar("T", abas(5, 20), DIA2))), 1)
        sem_config = abas(5, 36)
        del sem_config["_Config"]
        self.assertEqual(len(b.encolheu(antes, b.montar("T", sem_config, DIA2))), 1)


class Pasta(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name) / "cofre"
        self.addCleanup(self._tmp.cleanup)

    def arquivos(self, raiz=None):
        raiz = raiz or self.pasta
        return sorted(p.relative_to(raiz).as_posix() for p in raiz.rglob("*") if p.is_file())

    def test_grava_uma_vez_e_nao_repete_sem_mudanca(self):
        nome, gravou, sumiu = b.salvar_pasta(b.montar("T", abas(), DIA1), self.pasta)
        self.assertEqual((nome, gravou, sumiu), ("2026-10-01 07h40 backup cobrancas", True, []))
        nome2, gravou2, _ = b.salvar_pasta(b.montar("T", abas(), DIA2), self.pasta)
        self.assertEqual((nome2, gravou2), (nome, False))
        self.assertEqual(self.arquivos(), ["2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json",
                                           "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.xlsx", b.MARCA])
        self.assertIn("02/10/2026 às 13:00", (self.pasta / b.MARCA).read_text(encoding="utf-8"))

    def test_mudou_grava_outra_e_guarda_a_antiga(self):
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        antiga = (self.pasta / "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json").read_bytes()
        _, gravou, _ = b.salvar_pasta(b.montar("T", abas(3), DIA2), self.pasta)
        self.assertTrue(gravou)
        self.assertEqual(len(self.arquivos()), 5)
        self.assertEqual((self.pasta / "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json").read_bytes(), antiga)

    def test_cada_mes_na_sua_pasta_e_a_ultima_vale_entre_meses(self):
        novembro = datetime(2026, 11, 3, 7, 40, tzinfo=b.TIMEZONE)
        janeiro = datetime(2027, 1, 4, 7, 40, tzinfo=b.TIMEZONE)
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        b.salvar_pasta(b.montar("T", abas(3), novembro), self.pasta)
        b.salvar_pasta(b.montar("T", abas(4), janeiro), self.pasta)
        self.assertEqual([a for a in self.arquivos() if a.endswith(".xlsx")],
                         ["2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.xlsx",
                          "2026/11 - Novembro/2026-11-03 07h40 backup cobrancas.xlsx",
                          "2027/01 - Janeiro/2027-01-04 07h40 backup cobrancas.xlsx"])
        self.assertEqual(b.ultima_copia(self.pasta)["feito_em"], janeiro.isoformat(timespec="seconds"))
        _, gravou, _ = b.salvar_pasta(b.montar("T", abas(4), janeiro.replace(day=5)), self.pasta)
        self.assertFalse(gravou)

    def test_duas_copias_no_mesmo_minuto_a_nova_ganha_o_minuto_seguinte(self):
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        primeira = self.pasta / "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json"
        antes = primeira.read_bytes()
        nome, gravou, _ = b.salvar_pasta(b.montar("T", abas(3), DIA1.replace(second=50)), self.pasta)
        self.assertEqual((nome, gravou), ("2026-10-01 07h41 backup cobrancas", True))
        nome, gravou, _ = b.salvar_pasta(b.montar("T", abas(4), DIA1.replace(second=55)), self.pasta)
        self.assertEqual((nome, gravou), ("2026-10-01 07h42 backup cobrancas", True))
        self.assertEqual(primeira.read_bytes(), antes)
        self.assertEqual(len(b.ultima_copia(self.pasta)["abas"]["_Cobrancas"]), 5)   # a mais nova vale

    def test_espelho_aponta_arquivo_do_servidor_com_outro_tamanho_sem_mexer(self):
        destino = Path(self._tmp.name) / "servidor"
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        b.espelhar(self.pasta, destino)
        la = destino / "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.xlsx"
        la.write_bytes(la.read_bytes() + b"estragado")
        diferentes = []
        self.assertEqual(b.espelhar(self.pasta, destino, diferentes=diferentes), 0)
        self.assertEqual([Path(d).name for d in diferentes], ["2026-10-01 07h40 backup cobrancas.xlsx"])
        self.assertTrue(la.read_bytes().endswith(b"estragado"))
        limpo = []
        b.espelhar(self.pasta, Path(self._tmp.name) / "outro", diferentes=limpo)
        self.assertEqual(limpo, [])

    def test_espelho_nunca_reescreve_copia_datada(self):
        destino = Path(self._tmp.name) / "servidor"
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        b.espelhar(self.pasta, destino)
        la = destino / "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json"
        la.write_text("conteudo do servidor", encoding="utf-8")
        b.espelhar(self.pasta, destino)
        self.assertEqual(la.read_text(encoding="utf-8"), "conteudo do servidor")

    def test_cofre_guarda_so_o_json_e_dele_se_refaz_o_excel(self):
        copia = b.montar("T", abas(3), DIA1)
        nome, gravou, _ = b.salvar_pasta(copia, self.pasta, xlsx=False)
        self.assertEqual(self.arquivos(), ["2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json", b.MARCA])
        self.assertIn(nome + ".json", (self.pasta / b.MARCA).read_text(encoding="utf-8"))
        _, gravou, _ = b.salvar_pasta(b.montar("T", abas(3), DIA2), self.pasta, xlsx=False)
        self.assertFalse(gravou)
        guardado = json.loads((self.pasta / "2026/10 - Outubro" / (nome + ".json")).read_text(encoding="utf-8"))
        self.assertEqual(b.ler_xlsx(b.gerar_xlsx(guardado)), copia["abas"])

    def test_copia_em_formato_antigo_e_refeita_uma_vez(self):
        velha = b.montar("T", abas(), DIA1)
        velha["formato"] = 1
        b.salvar_pasta(velha, self.pasta)
        _, gravou, _ = b.salvar_pasta(b.montar("T", abas(), DIA2), self.pasta)
        self.assertTrue(gravou)
        _, gravou, _ = b.salvar_pasta(b.montar("T", abas(), DIA2.replace(hour=19)), self.pasta)
        self.assertFalse(gravou)
        self.assertEqual(len(self.arquivos()), 5)

    def test_encolheu_guarda_a_copia_e_avisa(self):
        b.salvar_pasta(b.montar("T", abas(5), DIA1), self.pasta)
        _, gravou, sumiu = b.salvar_pasta(b.montar("T", abas(1), DIA2), self.pasta)
        self.assertTrue(gravou and sumiu)

    def test_json_guardado_reconstroi_a_planilha(self):
        copia = b.montar("T", abas(), DIA1)
        b.salvar_pasta(copia, self.pasta)
        lido = json.loads((self.pasta / "2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json").read_text(encoding="utf-8"))
        self.assertEqual(lido["abas"], copia["abas"])
        self.assertEqual(b.impressao_digital(lido["abas"]), lido["impressao"])

    def test_json_cortado_no_meio_e_ignorado(self):
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        (self.pasta / "2026/10 - Outubro/2026-10-05 09h00 backup cobrancas.json").write_text('{"abas": {', encoding="utf-8")
        self.assertEqual(b.ultima_copia(self.pasta)["feito_em"], DIA1.isoformat(timespec="seconds"))

    def test_espelho_leva_so_o_que_falta_e_nunca_apaga(self):
        destino = Path(self._tmp.name) / "servidor"
        b.salvar_pasta(b.montar("T", abas(2), DIA1), self.pasta)
        self.assertEqual(b.espelhar(self.pasta, destino), 3)
        self.assertEqual(b.espelhar(self.pasta, destino), 0)
        (destino / "so_no_servidor.txt").write_text("x", encoding="utf-8")
        b.salvar_pasta(b.montar("T", abas(3), DIA2), self.pasta)
        self.assertEqual(b.espelhar(self.pasta, destino), 3)   # 2 arquivos novos + a marca atualizada
        self.assertEqual(len(self.arquivos(destino)), 6)
        self.assertEqual((destino / b.MARCA).read_bytes(), (self.pasta / b.MARCA).read_bytes())


class Execucao(unittest.TestCase):
    def test_planilha_errada_nao_grava_nada(self):
        with tempfile.TemporaryDirectory() as tmp:
            erros = b.executar(b.montar("Outra", {"Página1": [["a"]]}, DIA1), Path(tmp) / "c", saida=lambda *_: None)
            self.assertTrue(erros)
            self.assertFalse((Path(tmp) / "c").exists())

    def test_espelho_fora_do_ar_nao_impede_a_pasta(self):
        with tempfile.TemporaryDirectory() as tmp:
            bloqueio = Path(tmp) / "arquivo"
            bloqueio.write_text("x", encoding="utf-8")   # um arquivo no lugar da pasta: o espelho falha
            erros = b.executar(b.montar("T", abas(), DIA1), Path(tmp) / "c", bloqueio / "dentro",
                               saida=lambda *_: None)
            self.assertEqual(len(erros), 1)
            self.assertTrue((Path(tmp) / "c/2026/10 - Outubro/2026-10-01 07h40 backup cobrancas.json").exists())

    def test_drive_sem_chave_e_pulado_sem_erro(self):
        import os
        os.environ.pop("DRIVE_OAUTH", None)
        with tempfile.TemporaryDirectory() as tmp:
            falas = []
            self.assertEqual(b.executar(b.montar("T", abas(), DIA1), Path(tmp) / "c", drive=True,
                                        saida=falas.append), [])
            self.assertTrue(any("pulando" in f for f in falas))


if __name__ == "__main__":
    unittest.main()
